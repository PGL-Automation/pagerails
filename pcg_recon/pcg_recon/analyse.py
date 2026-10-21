"""Bank-side statement analysis — used to onboard a new bank / account before the
book side is connected, and as a monthly health check of the statement itself.

    pcg-recon analyse --file STATEMENT.xls --profile providus_xls --account PAM-PRV-CLIENT

Produces a workbook: control-total tie-out, running-balance continuity, category
mix by month (live SUMIFS over the normalised lines), reversal pairing, bank-charge
audit, flows with Page Group entities, and every line normalised and classified.
"""
from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .charges import audit, session_of
from .ingest.tabular import read_tabular
from .matcher import Matcher
from .model import fmt
from .reconcile import classify
from .report import BODY, BOLD, NUM, TITLE, F, GREEN, RED, AMBER, _table

GROUP_NAMES = re.compile(r"PAGE ASSET MANAGEMENT LIMITED|PAGE ASSET MGT(?! - +CLIENTS)|PAGE CAPITAL LIMITED|"
                         r"PAGE INTERNATIONAL FINANCE|PAGE FINANCIALS", re.I)
_TYPE = re.compile(r"^([A-Z][A-Z/().&\\ -]*?)(?=\s+\d|\s{2,}|\s*$)")
_OUT_CP = re.compile(r"\bTO\s+([A-Z0-9 &.\-]+?)\s*\|\s*([^|/]+)", re.I)
_IN_CP = re.compile(r"\bFROM\s+([A-Z0-9 &.\-]+?)\s*/\s*([^/|]+?)(?:-|/|$)", re.I)


def txn_type(desc: str) -> str:
    m = _TYPE.match(desc.strip().upper())
    return (m.group(1).strip() if m else desc.strip().upper()[:30]) or "?"


_BANKS = [("FBNQUEST", "FBNQUEST MERCHANT BANK"), ("FIRST CITY", "FCMB"), ("FCMB", "FCMB"),
          ("FIRST BANK", "FIRST BANK"), ("FBN", "FIRST BANK"), ("UNITED BANK FOR AFRICA", "UBA"), ("UBA", "UBA"),
          ("WEMA", "WEMA"), ("ABBEY", "ABBEY MORTGAGE BANK"), ("RAND MERCHANT", "RAND MERCHANT BANK"),
          ("STANBIC", "STANBIC IBTC"), ("GTB", "GTBANK"), ("GUARANTY", "GTBANK"), ("ACCESS", "ACCESS"),
          ("ZENITH", "ZENITH"), ("UNION", "UNION"), ("STERLING", "STERLING"), ("FIDELITY", "FIDELITY"),
          ("POLARIS", "POLARIS"), ("SUNTRUST", "SUNTRUST"), ("FSDH", "FSDH"), ("STANDARD", "STANDARD CHARTERED"),
          ("KEYSTONE", "KEYSTONE"), ("PARALLEX", "PARALLEX"), ("GLOBUS", "GLOBUS"), ("ECOBANK", "ECOBANK"),
          ("PROVIDUS", "PROVIDUS"), ("UNDEFINED", "PROVIDUS (internal)")]


def norm_bank(b: str) -> str:
    u = b.upper()
    for k, v in _BANKS:
        if k in u:
            return v
    return u


def counterparty(desc: str, amount: int) -> tuple[str, str]:
    m = (_OUT_CP if amount < 0 else _IN_CP).search(desc)
    if not m:
        m = (_IN_CP if amount < 0 else _OUT_CP).search(desc)
    if not m:
        return "", ""
    bank = re.sub(r"^TRZL-", "", m.group(1).strip(), flags=re.I)
    return norm_bank(bank)[:30], re.sub(r"\s+", " ", m.group(2)).strip()[:60]


def analyse(path: str, profile: dict, rules: dict, account_id: str, out: str, account_label: str = ""):
    txns, meta = read_tabular(path, profile, account_id, "BANK")
    m = Matcher(txns, [], rules, 3)
    m.p0_reversals()
    rev_pairs = m.matches
    in_rev = {u: mt for mt in rev_pairs for u in mt.bank_uids}
    charge_rows, charge_summary = audit(txns, rules)

    rows = []
    for t in txns:
        cat = classify(t, rules)[0]
        bank, cp = counterparty(t.description, t.amount)
        grp = bool(GROUP_NAMES.search(t.description))
        mt = in_rev.get(t.uid)
        rows.append(dict(date=t.txn_date, month=t.txn_date.strftime("%Y-%m"), type=txn_type(t.description),
                         category=cat, debit=-t.amount / 100 if t.amount < 0 else None,
                         credit=t.amount / 100 if t.amount > 0 else None, balance=t.running_balance / 100,
                         session=(session_of(t) or "")[4:], doc=t.reference, cp_bank=bank, counterparty=cp,
                         group="YES" if grp else "", reversal=(mt.match_id + (" (review)" if mt.status == "PROPOSED" else "")) if mt else "",
                         narration=t.description, uid=t.uid, net=t.amount))

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    # ---------------- Transactions (built first; other sheets reference it)
    wt = wb.create_sheet("Transactions")
    th = ["Date", "Month", "Txn type (bank code)", "Category", "Debit", "Credit", "Balance", "NIP session",
          "Doc no", "Counterparty bank", "Counterparty (parsed)", "Page Group", "Reversal pair", "Narration", "UID"]
    trows = [[r["date"], r["month"], r["type"], r["category"], r["debit"], r["credit"], r["balance"], r["session"],
              r["doc"], r["cp_bank"], r["counterparty"], r["group"], r["reversal"], r["narration"], r["uid"]] for r in rows]
    _table(wt, th, trows, [11, 8, 30, 22, 17, 17, 17, 34, 9, 18, 30, 7, 12, 90, 20],
           num_cols=("Debit", "Credit", "Balance"))
    for i in range(2, len(trows) + 2):
        wt.cell(i, 1).number_format = "dd-mmm-yyyy"
    n = len(trows) + 1
    TB, TD, TE, TF, TL = (f"Transactions!$B$2:$B${n}", f"Transactions!$D$2:$D${n}", f"Transactions!$E$2:$E${n}",
                          f"Transactions!$F$2:$F${n}", f"Transactions!$L$2:$L${n}")

    # ---------------- Monthly by category (formulas)
    wm = wb.create_sheet("Monthly")
    months = sorted({r["month"] for r in rows})
    cats = [c for c, _ in Counter(r["category"] for r in rows).most_common()]
    wm["A1"] = "Lines, debits and credits by month and category — live SUMIFS / COUNTIFS over the Transactions sheet"
    wm["A1"].font = Font(name=F, size=9, italic=True)
    hdr = ["Category", "Measure"] + months + ["Total"]
    grid = []
    for c in cats:
        for meas in ("Lines", "Debits", "Credits"):
            i = len(grid) + 4
            row = [c if meas == "Lines" else "", meas]
            for j, mo in enumerate(months):
                if meas == "Lines":
                    row.append(f'=COUNTIFS({TD},"{c}",{TB},"{mo}")')
                else:
                    rng = TE if meas == "Debits" else TF
                    row.append(f'=SUMIFS({rng},{TD},"{c}",{TB},"{mo}")')
            row.append(f"=SUM(C{i}:{get_column_letter(2 + len(months))}{i})")
            grid.append(row)
    _table(wm, hdr, grid, [26, 9] + [16] * (len(months) + 1), start_row=3)
    for i, row in enumerate(grid, 4):
        for j in range(3, 4 + len(months)):
            wm.cell(i, j).number_format = "#,##0" if row[1] == "Lines" else NUM

    # ---------------- Reversals
    wr = wb.create_sheet("Reversals")
    by_uid = {t.uid: t for t in txns}
    rh = ["Pair", "Status", "Rule note", "Debit date", "Debit leg(s)", "Credit date", "Reversal credit", "Difference", "Narration"]
    rrows = []
    for mt in rev_pairs:
        legs = [by_uid[u] for u in mt.bank_uids]
        d = [x for x in legs if x.amount < 0]
        c = [x for x in legs if x.amount > 0]
        rrows.append([mt.match_id, mt.status, mt.note, min(x.txn_date for x in d), -sum(x.amount for x in d) / 100,
                      min(x.txn_date for x in c), sum(x.amount for x in c) / 100, mt.difference / 100,
                      d[0].description[:110]])
    rrows.sort(key=lambda r: (r[1] != "PROPOSED", r[3]))
    _table(wr, rh, rrows, [8, 10, 36, 11, 17, 11, 17, 11, 90], num_cols=("Debit leg(s)", "Reversal credit", "Difference"))
    for i, r in enumerate(rrows, 2):
        for j in (4, 6):
            wr.cell(i, j).number_format = "dd-mmm-yyyy"
        if r[1] == "PROPOSED":
            for j in range(1, 10):
                wr.cell(i, j).fill = AMBER
    unpaired_rev = [t for t in txns if t.amount > 0 and classify(t, rules)[0] == "REVERSAL" and t.uid not in in_rev]

    # ---------------- Charge audit
    wc = wb.create_sheet("Charge Audit")
    wc["A1"] = "Tariff observed per charged transfer (legs tied by NIP session ID)"
    wc["A1"].font = BOLD
    _table(wc, ["Commission", "VAT", "Stamp duty", "Transfers"], [list(x) for x in charge_summary["tariff"]],
           [14, 10, 12, 11], start_row=2, num_cols=("Commission", "VAT", "Stamp duty"))
    r0 = 4 + len(charge_summary["tariff"])
    wc.cell(r0, 1, "Exceptions").font = BOLD
    ch = ["Check", "Detail", "Amount at issue", "Date", "Direction", "Principal", "Commission", "VAT", "Stamp duty",
          "Reversed", "NIP session", "Narration"]
    crow = [[x["check"], x["detail"], x["amount_at_issue"], x["date"], x["direction"], x["principal"], x["commission"],
             x["vat"], x["stamp_duty"], x["reversed"], x["session"], x["narration"]] for x in charge_rows]
    _table(wc, ch, crow, [18, 44, 14, 11, 9, 17, 11, 8, 10, 17, 34, 80], start_row=r0 + 1,
           num_cols=("Amount at issue", "Principal", "Commission", "VAT", "Stamp duty", "Reversed"))
    wc.freeze_panes = None

    # ---------------- Group flows
    wg = wb.create_sheet("Group Flows")
    g = defaultdict(lambda: [0, 0, 0, 0])
    for r in rows:
        if r["group"] and r["category"] not in ("BANK_CHARGES", "VAT_ON_CHARGES", "EMTL_STAMP_DUTY"):
            name = GROUP_NAMES.search(r["narration"]).group(0).upper()
            name = "PAGE ASSET MANAGEMENT LIMITED" if name.startswith("PAGE ASSET") else name
            k = (name, r["cp_bank"] or "(own/internal)")
            if r["net"] < 0:
                g[k][0] += 1; g[k][1] += -r["net"]
            else:
                g[k][2] += 1; g[k][3] += r["net"]
    gh = ["Group entity named in narration", "Counterparty bank", "Out lines", "Out value", "In lines", "In value", "Net out"]
    grows = [[k[0], k[1], v[0], v[1] / 100, v[2], v[3] / 100, (v[1] - v[3]) / 100]
             for k, v in sorted(g.items(), key=lambda kv: -(kv[1][1] + kv[1][3]))]
    _table(wg, gh, grows, [36, 26, 10, 19, 9, 19, 19], num_cols=("Out value", "In value", "Net out"))

    # ---------------- Summary
    ctl = meta.__dict__.get("control_check", {})
    ws["A1"] = f"Statement analysis — {account_label or account_id}"
    ws["A1"].font = TITLE
    ws["A2"] = f"{Path(path).name} · {meta.period_start:%d %b %Y} to {meta.period_end:%d %b %Y} · generated {datetime.now():%d %b %Y %H:%M}"
    ws["A2"].font = Font(name=F, size=9, italic=True)
    ws.column_dimensions["A"].width = 46
    for col, w in zip("BCDE", (22, 22, 12, 60)):
        ws.column_dimensions[col].width = w
    r = 4
    ws.cell(r, 1, "1. Did every line come through?").font = BOLD
    r += 1
    for j, h in enumerate(["Control figure", "Printed by bank", "Parsed by engine", "Agrees"], 1):
        c = ws.cell(r, j, h); c.font = BOLD
    labels = {"opening": "Opening balance", "closing": "Closing balance", "total_debit": "Total debits",
              "total_credit": "Total credits", "count_debit": "Debit lines", "count_credit": "Credit lines"}
    for k, lab in labels.items():
        if k in ctl:
            r += 1
            bank_v, got, ok = ctl[k]
            ws.cell(r, 1, lab).font = BODY
            for j, v in ((2, bank_v), (3, got)):
                c = ws.cell(r, j, v if "count" in k else v / 100)
                c.number_format = "#,##0" if "count" in k else NUM
                c.font = BODY
            c = ws.cell(r, 4, "YES" if ok else "NO"); c.fill = GREEN if ok else RED; c.font = BOLD
    r += 1
    ws.cell(r, 1, "Running-balance breaks").font = BODY
    c = ws.cell(r, 3, len(meta.continuity_breaks)); c.font = BODY
    c = ws.cell(r, 4, "YES" if not meta.continuity_breaks else "NO"); c.fill = GREEN if not meta.continuity_breaks else RED; c.font = BOLD

    r += 2
    ws.cell(r, 1, "2. What is on the statement").font = BOLD
    cnt = Counter(x["category"] for x in rows)
    val = defaultdict(int)
    for x in rows:
        val[x["category"]] += abs(x["net"])
    charge_lines = sum(cnt[c] for c in ("BANK_CHARGES", "VAT_ON_CHARGES", "EMTL_STAMP_DUTY"))
    facts = [
        ("Transaction lines", f"{len(rows):,}"),
        ("Charge lines (commission, VAT, stamp duty)", f"{charge_lines:,}  ({charge_lines / len(rows):.0%} of lines)  ·  NGN {sum(val[c] for c in ('BANK_CHARGES', 'VAT_ON_CHARGES', 'EMTL_STAMP_DUTY')) / 100:,.2f}"),
        ("Lines carrying a NIP session ID", f"{sum(1 for x in rows if x['session']) / len(rows):.1%}"),
        ("Reversal lines (credits and corrections)", f"{cnt['REVERSAL']:,}  ·  reversal credits paired to their debits: {sum(1 for t in txns if t.amount > 0 and t.uid in in_rev and in_rev[t.uid].status == 'AUTO'):,}; for review: {sum(1 for mt in rev_pairs if mt.status == 'PROPOSED')}; unpaired: {len(unpaired_rev)}"),
        ("Time deposits placed from this account", f"{cnt['TIME_DEPOSIT_PLACED']} · NGN {val['TIME_DEPOSIT_PLACED'] / 100:,.2f}"),
        ("Time deposits returned to this account", f"{cnt['TIME_DEPOSIT_RETURNED']} · NGN {val['TIME_DEPOSIT_RETURNED'] / 100:,.2f}"),
        ("Charged transfers audited", f"{charge_summary.get('sessions_with_charges', 0):,}  ·  exceptions: {len(charge_rows)}"),
        ("Lines naming a Page Group entity (excl. charges)", f"{sum(v[0] + v[2] for v in g.values()):,}"),
    ]
    for k, v in facts:
        r += 1
        ws.cell(r, 1, k).font = BODY
        ws.cell(r, 2, v).font = BODY
    r += 2
    ws.cell(r, 1, "Sheets").font = BOLD
    for k, v in [("Monthly", "category mix by month (formulas)"), ("Reversals", "each failed transfer and its refund"),
                 ("Charge Audit", "tariff and per-transfer exceptions"), ("Group Flows", "flows naming Page Group entities"),
                 ("Transactions", "every line normalised, classified and referenced")]:
        r += 1
        ws.cell(r, 1, "   " + k).font = BODY
        ws.cell(r, 2, v).font = BODY

    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.move_sheet("Transactions", offset=len(wb.sheetnames))
    wb.save(out)
    return {"lines": len(rows), "control": ctl, "breaks": meta.continuity_breaks, "categories": cnt,
            "values": val, "reversals": rev_pairs, "unpaired_reversals": unpaired_rev,
            "charge_rows": charge_rows, "charge_summary": charge_summary, "group": g, "meta": meta}
