"""pagerails command line.

  pagerails run      --inputs DIR --period 2026-09 --maker "A. Preparer"   [--accounts PAM-OPS-01,...]
  pagerails decide   --file out/recon_pack_2026-09.xlsx                     (loads Review Queue decisions)
  pagerails signoff  --run RUN_ID --checker "B. Reviewer" [--note "..."]
  pagerails status   [--period 2026-09-30]
  pagerails verify-audit
"""
from __future__ import annotations

import argparse
import calendar
import hashlib
import sys
from datetime import date
from pathlib import Path

import yaml

from .engine import run_account
from .model import fmt
from .report import write_pack
from .store import Store

HERE = Path(__file__).resolve().parent.parent


def parse_period(s: str) -> tuple[date, date]:
    if ":" in s:
        a, b = s.split(":")
        return date.fromisoformat(a), date.fromisoformat(b)
    parts = s.split("-")
    if len(parts) == 2:
        y, m = int(parts[0]), int(parts[1])
        return date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])
    d = date.fromisoformat(s)
    return d, d


def load_config(cfg_dir: Path):
    accounts = yaml.safe_load((cfg_dir / "accounts.yaml").read_text())
    profiles = yaml.safe_load((cfg_dir / "profiles.yaml").read_text())
    rules_txt = (cfg_dir / "rules.yaml").read_text()
    rules = yaml.safe_load(rules_txt)
    return accounts, profiles, rules, hashlib.sha256(rules_txt.encode()).hexdigest()[:12]


def cmd_run(a):
    accounts, profiles, rules, rh = load_config(Path(a.config))
    period = parse_period(a.period)
    store = Store(a.db)
    wanted = set(a.accounts.split(",")) if a.accounts else None
    results, failures = [], []
    for acct in accounts["accounts"]:
        if wanted and acct["id"] not in wanted:
            continue
        if acct.get("enabled") is False and not (wanted and acct["id"] in wanted):
            continue
        try:
            res = run_account(acct, period, Path(a.inputs), store, rules, profiles, rh, a.maker)
            results.append(res)
            print(f"  {acct['id']:<14} bank {fmt(res.bank_close):>18}  book {fmt(res.book_close):>18}  "
                  f"diff {fmt(res.difference):>12}  auto {res.stats['auto_match_rate']:.0%}  "
                  f"open {len(res.open_bank) + len(res.open_book):>3}  {res.status}")
        except Exception as e:  # one bad feed must not stop the other accounts
            failures.append((acct["id"], str(e)))
            print(f"  {acct['id']:<14} FAILED: {e}", file=sys.stderr)
            store.audit(a.maker, "RUN_FAILED", {"account_id": acct["id"], "error": str(e)})
    if results:
        Path(a.out).mkdir(parents=True, exist_ok=True)
        label = a.period if len(a.period) != 10 else period[1].strftime("%d %b %Y")
        out = Path(a.out) / f"recon_pack_{a.period.replace(':', '_')}.xlsx"
        write_pack(results, str(out), label, a.maker, accounts["entities"])
        print(f"\nPack: {out}")
    return 1 if failures or any(r.difference for r in results) else 0


def cmd_decide(a):
    from openpyxl import load_workbook
    store = Store(a.db)
    ws = load_workbook(a.file, data_only=True)["Review Queue"]
    n = 0
    hdr = [c.value for c in ws[3]]
    for row in ws.iter_rows(min_row=4, values_only=True):
        d = dict(zip(hdr, row))
        act = (d.get("Decision") or "").strip().upper()
        if act not in ("APPROVE", "REJECT"):
            continue
        if not d.get("Reviewer"):
            print(f"  skip {d['Match ID']}: decision without reviewer name", file=sys.stderr)
            continue
        store.add_decision(d["Account"], act, (d["Bank UIDs"] or "").split(";"), (d["Book UIDs"] or "").split(";"),
                           d["Reviewer"], d.get("Note") or "")
        n += 1
    print(f"Recorded {n} decisions. Re-run the period to apply them.")
    return 0


def cmd_manual(a):
    store = Store(a.db)
    store.add_decision(a.account, "MANUAL", a.bank.split(",") if a.bank else [],
                       a.book.split(",") if a.book else [], a.reviewer, a.note or "")
    print("Manual match recorded. Re-run the period to apply it.")
    return 0


def cmd_signoff(a):
    Store(a.db).signoff(a.run, a.checker, a.note or "")
    print(f"Signed off {a.run} by {a.checker}")
    return 0


def cmd_analyse(a):
    from .analyse import analyse
    accounts, profiles, rules, rh = load_config(Path(a.config))
    Path(a.out).mkdir(parents=True, exist_ok=True)
    out = Path(a.out) / f"statement_analysis_{a.account}.xlsx"
    res = analyse(a.file, profiles[a.profile], rules, a.account, str(out), a.label or a.account)
    ok = all(v[2] for v in res["control"].values()) and not res["breaks"]
    print(f"{res['lines']:,} lines · control totals {'AGREE' if ok else 'DO NOT AGREE'} · "
          f"reversals paired {sum(1 for m in res['reversals'] if m.status == 'AUTO')} · "
          f"charge exceptions {len(res['charge_rows'])}\nWorkbook: {out}")
    return 0 if ok else 1


def cmd_status(a):
    for r in Store(a.db).runs(a.period):
        print(f"{r['run_id']:<34} {r['status']:<28} diff {fmt(r['difference']):>12}  "
              f"maker {r['maker']:<14} checker {r['checker'] or '-':<14} {r['signed_at'] or 'NOT SIGNED'}")
    return 0


def cmd_verify(a):
    ok, n = Store(a.db).verify_audit_chain()
    print(f"Audit chain {'INTACT' if ok else 'BROKEN at seq ' + str(n)} ({n} entries checked)")
    return 0 if ok else 2


def main(argv=None):
    p = argparse.ArgumentParser(prog="pagerails", description="PageRails — Financial Operations Platform")
    p.add_argument("--config", default=str(HERE / "config"))
    p.add_argument("--db", default="recon.db")
    sub = p.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--inputs", required=True)
    r.add_argument("--period", required=True, help="2026-09 | 2026-09-29 | 2026-09-01:2026-09-15")
    r.add_argument("--maker", required=True)
    r.add_argument("--accounts")
    r.add_argument("--out", default="out")
    d = sub.add_parser("decide")
    d.add_argument("--file", required=True)
    m = sub.add_parser("manual")
    m.add_argument("--account", required=True)
    m.add_argument("--bank", default="")
    m.add_argument("--book", default="")
    m.add_argument("--reviewer", required=True)
    m.add_argument("--note")
    s = sub.add_parser("signoff")
    s.add_argument("--run", required=True)
    s.add_argument("--checker", required=True)
    s.add_argument("--note")
    st = sub.add_parser("status")
    st.add_argument("--period")
    sub.add_parser("verify-audit")
    an = sub.add_parser("analyse", help="bank-side statement analysis (onboarding / health check)")
    an.add_argument("--file", required=True)
    an.add_argument("--profile", required=True)
    an.add_argument("--account", required=True)
    an.add_argument("--label")
    an.add_argument("--out", default="out")
    a = p.parse_args(argv)
    try:
        return {"run": cmd_run, "decide": cmd_decide, "manual": cmd_manual, "signoff": cmd_signoff,
                "status": cmd_status, "verify-audit": cmd_verify, "analyse": cmd_analyse}[a.cmd](a)
    except (PermissionError, ValueError, FileNotFoundError) as e:
        print(f"REFUSED: {e}", file=sys.stderr)
        return 3


if __name__ == "__main__":
    sys.exit(main())
