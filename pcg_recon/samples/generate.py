"""Synthetic test data: two months (Sep, Oct 2026) for all five example accounts,
in every supported statement format, with the known reconciling items a real
month throws up. Used by tests/ as the regression fixture.

Scenario per account (September):
  * normal matched flows with 0-2 day bank/book date lags
  * deposit in transit (booked 30 Sep, credited by bank 1 Oct)
  * unpresented cheque (issued Sep, presented Oct)
  * bank-only: EMTL, SMS, VAT, COT/maintenance, interest (+ WHT) → draft journals
  * bulk payment: 3 book lines settled as 1 bank debit            (P3 aggregate)
  * tranche receipt: 1 book receipt credited as 2 bank lines       (P3 aggregate)
  * failed transfer + bank reversal                                (P0 reversal)
  * client-money accounts: one unidentified credit                 (AML alert)
October clears the carried items and books the September charges.
"""
from __future__ import annotations

import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

OPEN = {"PAM-OPS-01": 48_250_000_00, "PAM-CLIENT-01": 312_400_000_00, "PAM-MMF-CUST": 1_850_000_000_00,
        "PCL-OPS-01": 96_700_000_00, "PCL-COLL-01": 27_300_000_00}
NAMES = ["ADEBAYO OLUWASEUN", "CHIOMA OKAFOR", "IBRAHIM MUSA", "FUNMILAYO ADEYEMI", "EMEKA NWOSU",
         "HALIMA BELLO", "TUNDE BAKARE", "NGOZI EZE", "YUSUF ABDULLAHI", "BOLANLE OGUNDIPE"]
VENDORS = ["LAGOS ENERGY LTD", "IKOYI PROPERTIES LTD", "DATASAFE NIGERIA LTD", "PRIMEAUDIT PARTNERS",
           "SWIFTNET SOLUTIONS", "ZENITH OFFICE SUPPLIES"]


def nip(rng):
    return "".join(rng.choice("0123456789") for _ in range(30))


def ev(db, dk, amt, bdesc, kdesc, bref="", kref="", dtype="", bank=True, book=True):
    return dict(db=db, dk=dk, amt=amt, bdesc=bdesc, kdesc=kdesc, bref=bref, kref=kref, dtype=dtype,
                bank=bank, book=book)


def events(acct: str, y: int, m: int, rng: random.Random, client: bool, carry: bool) -> list[dict]:
    E = []
    d0 = date(y, m, 1)
    last = date(y, m, 30) if m in (9, 11) else date(y, m, 31)

    def day(a, b):
        return d0 + timedelta(days=rng.randint(a, b))

    # ---- normal inflows
    for i in range(14 if client else 8):
        d = day(0, 27)
        amt = rng.randint(50, 5000) * 10_000_00 // 100 * 100
        who = rng.choice(NAMES)
        code = f"PAM-SUB-{rng.randint(100000, 999999)}" if acct.startswith("PAM") else f"PCL-LN-{rng.randint(10000, 99999)}"
        s = nip(rng)
        lag = rng.choice([0, 0, 0, 1, 2])
        if rng.random() < 0.6:
            E.append(ev(d, max(d - timedelta(days=lag), d0), amt, f"NIP TRF FROM {who}/{code}/{s}",
                        f"Receipt {who.title()} {code}", kref=code, dtype="RCPT"))
        else:  # book has no usable reference → amount/date matching
            E.append(ev(d, max(d - timedelta(days=lag), d0), amt, f"TRANSFER FROM {who}", f"Being amount received from {who.title()}",
                        dtype="RCPT"))
    # ---- normal outflows
    for i in range(10):
        d = day(0, 27)
        amt = -rng.randint(20, 3000) * 10_000_00 // 100 * 100
        v = rng.choice(VENDORS)
        pv = f"PV{m:02d}{rng.randint(1000, 9999)}"
        E.append(ev(d + timedelta(days=rng.choice([0, 1])), d, amt, f"TRF TO {v} {pv} REF:FT26{rng.randint(10**7, 10**8)}",
                    f"Payment to {v.title()} {pv}", kref=pv, dtype="PMT"))
    # ---- two same-amount payments to the same vendor (ambiguity)
    d = day(5, 10)
    for k in range(2):
        E.append(ev(d + timedelta(days=k), d + timedelta(days=k), -1_500_000_00, f"TRF TO SWIFTNET SOLUTIONS MONTHLY RETAINER",
                    "Swiftnet retainer", dtype="PMT"))
    # ---- bulk payment: 3 book lines, 1 bank line
    d = day(20, 24)
    parts = [-4_125_000_00, -2_870_000_00, -1_655_500_00]
    batch = f"BULK{m:02d}{rng.randint(100, 999)}"
    E.append(ev(d, None, sum(parts), f"BULK PAYMENT BATCH {batch} 3 ITEMS", "", book=False))
    for i, p in enumerate(parts):
        E.append(ev(None, d, p, "", f"Bulk batch {batch} item {i + 1}", kref=batch, dtype="PMT", bank=False))
    # ---- tranche receipt: 1 book, 2 bank
    d = day(12, 16)
    E.append(ev(None, d, 20_000_000_00, "", "Receipt Afolabi Holdings settlement", kref="AFH-2026-091", bank=False))
    E.append(ev(d, None, 12_000_000_00, "TRF FROM AFOLABI HOLDINGS PART 1", "", book=False))
    E.append(ev(d + timedelta(days=1), None, 8_000_000_00, "TRF FROM AFOLABI HOLDINGS PART 2", "", book=False))
    # ---- failed transfer + reversal (bank only, net zero)
    d = day(8, 18)
    s = nip(rng)
    E.append(ev(d, None, -3_450_000_00, f"NIP TRF TO EKO GLOBAL VENTURES {s}", "", book=False))
    E.append(ev(d + timedelta(days=1), None, 3_450_000_00, f"REVERSAL NIP TRF TO EKO GLOBAL VENTURES {s}", "", book=False))
    # ---- cheque issued and presented within month
    d = day(2, 6)
    chq = rng.randint(100, 900)
    E.append(ev(d + timedelta(days=6), d, -2_200_000_00, f"CLG CHQ {chq:06d} PRESENTED", f"Chq {chq:06d} Ikeja Legal Chambers",
                kref=f"CHQ {chq:06d}", dtype="CHQ"))
    # ---- bank-only charges (booked by finance NEXT month)
    for k in range(rng.randint(6, 10)):
        E.append(ev(day(0, 29), None, -50_00, "EMTL CHARGE", "", book=False))
    E.append(ev(last, None, -4_000_00, f"SMS ALERT CHARGES {d0:%b %Y}".upper(), "", book=False))
    E.append(ev(last, None, -300_00, "VAT ON SMS ALERT", "", book=False))
    if not client:
        E.append(ev(last, None, -12_500_00, f"ACCOUNT MAINTENANCE FEE {d0:%b%y}".upper(), "", book=False))
        E.append(ev(last, None, -937_50, "VAT ON ACCOUNT MAINTENANCE FEE", "", book=False))
    if acct in ("PAM-MMF-CUST", "PCL-OPS-01"):
        E.append(ev(last, None, 1_845_210_45, "INTEREST CREDIT CALL DEPOSIT", "", book=False))
        E.append(ev(last, None, -184_521_05, "WHT ON INTEREST", "", book=False))
    # ---- client money: unidentified receipt
    if client:
        E.append(ev(day(20, 25), None, 7_500_000_00, f"NIP TRF FROM OKONKWO C/{nip(rng)}", "", book=False))
    # ---- year-end-style timing items
    if not carry:
        E.append(ev(None, last, 9_800_000_00, "", "Receipt Kolawole Estate - cheque lodged", kref="LODG-0930", bank=False))
        chq2 = rng.randint(901, 999)
        E.append(ev(None, last - timedelta(days=4), -6_350_000_00, "", f"Chq {chq2:06d} Victoria Island Facilities",
                    kref=f"CHQ {chq2:06d}", dtype="CHQ", bank=False))
    return E


def october_clearances(sep_events, acct):
    """Carried items clear in Oct; Sep bank-only charges get booked by finance."""
    E = []
    groups = {}
    for e in sep_events:
        if e["book"] and not e["bank"] and e["kref"] == "LODG-0930":
            E.append(ev(date(2026, 10, 1), None, e["amt"], "CHQ LODGEMENT KOLAWOLE ESTATE CLEARED", "", book=False))
        if e["book"] and not e["bank"] and e["dtype"] == "CHQ":
            E.append(ev(date(2026, 10, 5), None, e["amt"], f"CLG {e['kref']} PRESENTED", "", book=False))
        if e["bank"] and not e["book"]:
            for key, label in (("VAT", "VAT on bank charges"), ("EMTL", "EMTL levy"), ("SMS", "SMS alert charges"),
                               ("MAINTENANCE", "Account maintenance fee"), ("WHT", "WHT on interest"),
                               ("INTEREST", "Interest credit")):
                if key in e["bdesc"]:
                    groups.setdefault(label, 0)
                    groups[label] += e["amt"]
                    break
    # finance books one catch-up JV per charge type, early the next month
    for label, amt in groups.items():
        E.append(ev(None, date(2026, 10, 2), amt, "", f"Sep 2026 {label}", kref=f"JV-BC-{label[:4].upper()}",
                    dtype="JV", bank=False))
    return E


# ---------------------------------------------------------------- writers
def _bank_rows(E, opening):
    rows = sorted([e for e in E if e["bank"]], key=lambda e: e["db"])
    bal = opening
    out = []
    for e in rows:
        bal += e["amt"]
        out.append((e["db"], e["bdesc"], e["bref"], e["amt"], bal))
    return out, bal


def _book_rows(E, opening):
    rows = sorted([e for e in E if e["book"]], key=lambda e: e["dk"])
    bal = opening
    out = []
    for e in rows:
        bal += e["amt"]
        out.append((e["dk"], e["kdesc"], e["kref"], e["amt"], bal, e["dtype"]))
    return out, bal


def m(v):  # minor -> 2dp string
    return f"{v / 100:.2f}"


def write_all(root: Path, seed: int = 7):
    rng = random.Random(seed)
    truth = {}
    for acct, client in [("PAM-OPS-01", False), ("PAM-CLIENT-01", True), ("PAM-MMF-CUST", True),
                         ("PCL-OPS-01", False), ("PCL-COLL-01", False)]:
        d = root / acct
        d.mkdir(parents=True, exist_ok=True)
        sep = events(acct, 2026, 9, rng, client, carry=False)
        octo = events(acct, 2026, 10, rng, client, carry=True) + october_clearances(sep, acct)
        bank_open, book_open = OPEN[acct], OPEN[acct]
        truth[acct] = {}
        for label, E in (("2026-09", sep), ("2026-10", octo)):
            brows, bank_close = _bank_rows(E, bank_open)
            krows, book_close = _book_rows(E, book_open)
            truth[acct][label] = {"bank_close": bank_close, "book_close": book_close,
                                  "bank_open": bank_open, "book_open": book_open}
            WRITERS[acct](d, label, brows, krows, bank_open, book_open)
            bank_open, book_open = bank_close, book_close
    return truth


def w_pam_ops(d, label, brows, krows, bo, ko):
    with open(d / f"bank_{label}.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["BANK A PLC - ACCOUNT STATEMENT"])
        w.writerow(["Account Name:", "PAGE ASSET MANAGEMENT LIMITED"])
        w.writerow(["Account Number:", "****1101"])
        w.writerow([f"Period: {label}"])
        w.writerow([])
        w.writerow(["Trans Date", "Reference", "Value Date", "Narration", "Debit", "Credit", "Balance"])
        for dt, desc, ref, amt, bal in reversed(brows):          # portal exports newest-first
            w.writerow([dt.strftime("%d-%b-%Y"), ref, dt.strftime("%d-%b-%Y"), desc,
                        f"{-amt / 100:,.2f}" if amt < 0 else "", f"{amt / 100:,.2f}" if amt > 0 else "", f"{bal / 100:,.2f}"])
    with open(d / f"gl_{label}.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Date", "Type", "Reference", "Details", "Debit", "Credit"])
        for dt, desc, ref, amt, bal, ty in krows:
            w.writerow([dt.strftime("%d/%m/%Y"), ty, ref, desc, m(amt) if amt > 0 else "", m(-amt) if amt < 0 else ""])


def w_pam_client(d, label, brows, krows, bo, ko):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(d / f"statement_{label}.pdf"), pagesize=A4)
    y = 800

    def line(txt, size=8):
        nonlocal y
        if y < 50:
            c.showPage()
            y = 800
        c.setFont("Helvetica", size)
        c.drawString(30, y, txt)
        y -= 12

    line("BANK B PLC  -  STATEMENT OF ACCOUNT", 11)
    line("PAGE ASSET MANAGEMENT LTD - CLIENT COLLECTIONS  ****2202")
    line(f"Statement period {label}")
    line("Date        Value Date   Narration                                            Amount        Balance")
    line(f"OPENING BALANCE {bo / 100:,.2f}")
    for dt, desc, ref, amt, bal in brows:
        head, tail = desc[:48], desc[48:]
        line(f"{dt:%d-%b-%Y}  {dt:%d-%b-%Y}  {head:<50} {abs(amt) / 100:,.2f}  {bal / 100:,.2f}")
        if tail:
            line(f"                          {tail}")
    line(f"CLOSING BALANCE {brows[-1][4] / 100:,.2f}")
    c.save()
    wb = Workbook()
    ws = wb.active
    ws.append(["PAGE ASSET MANAGEMENT - CLIENT COLLECTIONS CASHBOOK"])
    ws.append([])
    ws.append(["Date", "Voucher No", "Particulars", "Receipts", "Payments", "Balance"])
    ws.append([None, "", "Balance b/f", None, None, ko / 100])
    for dt, desc, ref, amt, bal, ty in krows:
        ws.append([dt, ref, desc, amt / 100 if amt > 0 else None, -amt / 100 if amt < 0 else None, bal / 100])
    wb.save(d / f"cashbook_{label}.xlsx")


def w_mmf(d, label, brows, krows, bo, ko):
    with open(d / f"custodian_{label}.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["CUSTODIAN C - CASH STATEMENT", "PAGE MONEY MARKET FUND (EXAMPLE)"])
        w.writerow(["Settlement Date", "Transaction Description", "Trade Reference", "Net Amount", "Cash Balance"])
        for dt, desc, ref, amt, bal in brows:
            w.writerow([dt.isoformat(), desc, ref, m(amt), m(bal)])
    with open(d / f"fund_gl_{label}.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Date", "Reference", "Label", "Account", "Debit", "Credit"])
        for dt, desc, ref, amt, bal, ty in krows:
            w.writerow([dt.isoformat(), ref, desc, "F1-1001", m(amt) if amt > 0 else "0", m(-amt) if amt < 0 else "0"])
            w.writerow([dt.isoformat(), ref, desc, "F1-2001", m(-amt) if amt < 0 else "0", m(amt) if amt > 0 else "0"])


def w_pcl_ops(d, label, brows, krows, bo, ko):
    def bal(v, dt):
        return f"{'C' if v >= 0 else 'D'}{dt:%y%m%d}NGN{abs(v) // 100},{abs(v) % 100:02d}"
    L = ["{1:F01BANKDNGLAXXX0000000000}{2:I940PAGECAPXXXXN}{4:", ":20:STMT" + label.replace("-", ""),
         ":25:****4404", ":28C:00001/001", ":60F:" + bal(bo, brows[0][0])]
    for dt, desc, ref, amt, b in brows:
        mark = "C" if amt > 0 else "D"
        r = (ref or "NONREF")[:16]
        L.append(f":61:{dt:%y%m%d}{dt:%m%d}{mark}{abs(amt) // 100},{abs(amt) % 100:02d}NTRF{r}//{dt:%m%d}{abs(hash(desc)) % 10**8}")
        L.append(":86:" + desc[:65])
        if len(desc) > 65:
            L.append(desc[65:130])
    L.append(":62F:" + bal(brows[-1][4], brows[-1][0]))
    L.append("-}")
    (d / f"PCL_OPS_{label}.sta").write_text("\n".join(L) + "\n")
    with open(d / f"register_{label}.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Date", "Type", "Num", "Memo", "Payment", "Deposit"])
        for dt, desc, ref, amt, b, ty in krows:
            w.writerow([dt.strftime("%m/%d/%Y"), ty, ref, desc, m(-amt) if amt < 0 else "", m(amt) if amt > 0 else ""])


def w_pcl_coll(d, label, brows, krows, bo, ko):
    wb = Workbook()
    ws = wb.active
    ws.append(["BANK E - TRANSACTION HISTORY"])
    ws.append(["PAGE CAPITAL LIMITED - LOAN COLLECTIONS", "****5505"])
    ws.append([])
    ws.append(["Tran ID", "Value Date", "Narration", "Amount", "DR/CR", "Balance"])
    for i, (dt, desc, ref, amt, bal) in enumerate(brows):
        ws.append([f"S{dt:%y%m%d}{i:04d}", dt, desc, abs(amt) / 100, "D" if amt < 0 else "C", bal / 100])
    wb.save(d / f"bank_{label}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Date", "Voucher No", "Particulars", "Receipts", "Payments", "Balance"])
    ws.append([None, "", "Balance b/f", None, None, ko / 100])
    for dt, desc, ref, amt, bal, ty in krows:
        ws.append([dt, ref, desc, amt / 100 if amt > 0 else None, -amt / 100 if amt < 0 else None, bal / 100])
    wb.save(d / f"cashbook_{label}.xlsx")


WRITERS = {"PAM-OPS-01": w_pam_ops, "PAM-CLIENT-01": w_pam_client, "PAM-MMF-CUST": w_mmf,
           "PCL-OPS-01": w_pcl_ops, "PCL-COLL-01": w_pcl_coll}

if __name__ == "__main__":
    import json
    import sys
    root = Path(sys.argv[1] if len(sys.argv) > 1 else "samples/inputs")
    t = write_all(root)
    (root / "truth.json").write_text(json.dumps(t, indent=2))
    print(f"wrote samples to {root}")
